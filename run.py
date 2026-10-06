"""Run the whole analysis:  python run.py [--data data] [--out outputs]
Writes CSVs, an Excel board pack and qa_report.md. No network, no API keys, no paid calls."""
import argparse, json, sys
from pathlib import Path
import numpy as np, pandas as pd
from vireo.pipeline import run_all, DEFAULT_CODE

def lakh(x): return f"Rs {x/1e5:,.2f} lakh"

def headline(R):
    t, ex = R["tickets"], R["exceptions"]
    r = t[t.refund_inr.notna()]
    last4 = sorted(t.quarter.unique())[-4:]
    dd4 = ex[ex.quarter.isin(last4)]; r4 = r[r.quarter.isin(last4)]
    rate = len(dd4) / len(r4)
    target = 0.02
    q = len(last4)
    h = dict(
        total_refund_inr=float(r.refund_inr.sum()), refund_tickets=int(len(r)),
        q2_2026_inr=float(r[r.quarter == "2026Q2"].refund_inr.sum()),
        q2_2026_resolved_only_inr=float(r[(r.quarter == "2026Q2") & (r.status == "resolved")].refund_inr.sum()),
        double_dips_18m=int(len(ex)), double_dips_flagged=int((ex.double_dip == "flag").sum()),
        double_dips_notes_only=int((ex.double_dip == "notes only").sum()),
        double_dip_refund_value_18m=float(ex.refund_inr.sum()), leak_low_18m=float(ex.leak_low_inr.sum()),
        last4_quarters=last4, dd_rate_last4=rate, dd_count_last4=int(len(dd4)),
        leak_low_per_quarter_now=float(dd4.leak_low_inr.sum() / q),
        leak_high_per_quarter_now=float(dd4.leak_high_inr.sum() / q),
        target_rate=target,
        saving_low_per_quarter_at_target=float(dd4.leak_low_inr.sum() / q * (1 - target / rate)),
        saving_high_per_quarter_at_target=float(dd4.leak_high_inr.sum() / q * (1 - target / rate)),
        goodwill_over_cap_tickets=int(((r.reason_final == "GOODWILL") & (r.refund_inr > 500)).sum()),
        goodwill_over_cap_inr=float(r[(r.reason_final == "GOODWILL") & (r.refund_inr > 500)].refund_inr.sum()),
        filed_default_share_value=float(r[r.reason_filed == DEFAULT_CODE].refund_inr.sum() / r.refund_inr.sum()),
        true_goodwill_share_value=float(r[r.reason_final == "GOODWILL"].refund_inr.sum() / r.refund_inr.sum()),
        over_refunded_orders=int(len(R["over_refunded"])), over_refunded_inr=float(R["over_refunded"].over_by.sum()),
    )
    return h

def add_total_row(ws, first_data_row, last_data_row, first_col, last_col, label_col=1):
    from openpyxl.utils import get_column_letter as L
    tr = last_data_row + 1
    ws.cell(tr, label_col, "TOTAL")
    for c in range(first_col, last_col + 1):
        ws.cell(tr, c, f"=SUM({L(c)}{first_data_row}:{L(c)}{last_data_row})")
        ws.cell(tr, c).number_format = "#,##0"
    for c in range(1, last_col + 1):
        ws.cell(tr, c).font = ws.cell(tr, c).font.copy(bold=True)

def write_excel(R, h, path):
    s, a, t = R["summaries"], R["audit"], R["tickets"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        bridge = pd.DataFrame([
            ("Raw refund column, as exported (all rows)", a["raw_refund_sum"]),
            ("less: legacy rows that are re-import duplicates of a helpdesk row", -a["raw_refund_in_removed_dups"]),
            ("less: legacy unit overstatement (kept legacy rows x100)", -(a["legacy_raw_refund_kept"] - a["legacy_raw_refund_kept"] / 100)),
            ("= Reconciled refunds raised, Jan-25 to Jun-26", a["corrected_refund_sum"])], columns=["step", "inr"])
        bridge.to_excel(xw, sheet_name="1_Reconciliation", index=False)
        q = s["quarterly"].reset_index(); q.to_excel(xw, sheet_name="2_Quarterly", index=False)
        st = s["by_status"].reset_index(); st.to_excel(xw, sheet_name="2b_Quarter_by_status", index=False)
        s["monthly_by_reason_final"].reset_index().to_excel(xw, sheet_name="3_Month_x_Reason_FINAL", index=False)
        s["monthly_by_reason_filed"].reset_index().to_excel(xw, sheet_name="4_Month_x_Reason_AS_FILED", index=False)
        s["monthly_by_agent"].reset_index().to_excel(xw, sheet_name="5_Month_x_Agent", index=False)
        s["agents"].to_excel(xw, sheet_name="6_Agents", index=False)
        s["monthly_agent_reason"].to_excel(xw, sheet_name="7_Month_Agent_Reason_long", index=False)
        ex = R["exceptions"][["ticket_id", "created_at", "month", "agent_id", "agent_name", "agent_team", "product_sku", "refund_inr",
                              "reason_filed", "reason_final", "replacement_issued", "double_dip", "replacement_cost_inr", "leak_low_inr", "leak_high_inr", "agent_notes"]]
        ex.to_excel(xw, sheet_name="8_Refund_plus_unit", index=False)
        R["over_refunded"].reset_index().to_excel(xw, sheet_name="9_Orders_refunded_gt_value", index=False)
        r = t[t.refund_inr.notna()]
        r[["ticket_id", "created_at", "month", "status", "agent_id", "agent_name", "refund_inr", "reason_filed", "reason_final",
           "reason_source", "reason_conf", "source_system", "customer_message", "agent_notes"]].to_excel(xw, sheet_name="10_All_refund_tickets", index=False)
        pd.DataFrame([(k, json.dumps(v) if not isinstance(v, (int, float, str)) else v) for k, v in {**R["audit"], **R["recode_qa"]}.items()],
                     columns=["check", "value"]).to_excel(xw, sheet_name="11_QA", index=False)
        wb = xw.book
        for name, ncols_numeric_from in [("3_Month_x_Reason_FINAL", 2), ("4_Month_x_Reason_AS_FILED", 2)]:
            ws = wb[name]; add_total_row(ws, 2, ws.max_row, ncols_numeric_from, ws.max_column)
        ws = wb["5_Month_x_Agent"]; add_total_row(ws, 2, ws.max_row, 5, ws.max_column)
        ws = wb["6_Agents"]; add_total_row(ws, 2, ws.max_row, 8, 8)
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            for col in ws.columns:
                w = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:60])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, w + 2), 60)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, float): c.number_format = "#,##0"

def qa_report(R, h):
    a, q = R["audit"], R["recode_qa"]
    L = ["# QA report (generated by run.py)", "",
         "## Money reconciliation (deterministic - any failure here is a bug)", "",
         f"- Rows in export: {a['raw_rows']:,}; ticket_ids under both systems: {a['dup_ticket_ids']}; rows removed: {a['dup_rows_removed']}",
         f"- Duplicate pairs with money on both sides: {a['pairs_checked_for_unit']}; legacy = exactly 100x helpdesk in {a['pairs_exactly_x100']} of them",
         f"- Fields that differ inside a duplicate pair besides money: {a['pair_fields_that_differ_besides_money'] or 'none'}",
         f"- Raw refund column summed: Rs {a['raw_refund_sum']:,.0f}  ->  reconciled: Rs {a['corrected_refund_sum']:,.0f} ({a['refund_tickets']:,} refund tickets)",
         f"- Tickets with no roster match: {a['agents_unmatched_tickets']}", "",
         "## Reason re-coding (learned - this is where errors live)", "",
         f"- {q['default_code_rows']} refunds were filed under the dropdown default {DEFAULT_CODE} = {h['filed_default_share_value']:.1%} of value.",
         f"- Of those: {q['default_rows_goodwill_language']} show goodwill/giveaway language, {q['default_rows_recoded_by_model']} re-coded by text model, {q['default_rows_unclear']} left UNCLEAR.",
         f"- 5-fold CV on {q['labelled_rows']} refunds where the agent chose a deliberate code: accuracy {q['cv_accuracy']:.1%}; "
         f"{q['cv_accuracy_at_conf']:.1%} on the {q['cv_coverage_at_conf']:.0%} above confidence {q['min_conf']}.",
         "- Per class: " + ", ".join(f"{k} {v:.0%}" for k, v in q["cv_accuracy_per_class"].items()),
         "- Caveat: CV measures agreement with agents' deliberate codes, not truth, and those rows may be easier than the default-code rows.", ""]
    return "\n".join(L)

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--data", default="data"); ap.add_argument("--out", default="outputs")
    ns = ap.parse_args(); out = Path(ns.out); out.mkdir(exist_ok=True, parents=True)
    R = run_all(ns.data); h = headline(R); s = R["summaries"]
    s["monthly_by_reason_final"].to_csv(out / "monthly_by_reason_final.csv")
    s["monthly_by_reason_filed"].to_csv(out / "monthly_by_reason_as_filed.csv")
    s["monthly_by_agent"].to_csv(out / "monthly_by_agent.csv")
    s["monthly_agent_reason"].to_csv(out / "monthly_agent_reason_long.csv", index=False)
    s["agents"].to_csv(out / "agents.csv", index=False)
    s["quarterly"].to_csv(out / "quarterly.csv")
    R["exceptions"].to_csv(out / "refund_plus_unit_exceptions.csv", index=False)
    R["over_refunded"].to_csv(out / "orders_refunded_more_than_value.csv")
    write_excel(R, h, out / "vireo_refund_board_pack.xlsx")
    (out / "headline.json").write_text(json.dumps(h, indent=1))
    (out / "qa_report.md").write_text(qa_report(R, h))
    print(json.dumps(h, indent=1)); print("\nWrote", out.resolve())

if __name__ == "__main__":
    main()
